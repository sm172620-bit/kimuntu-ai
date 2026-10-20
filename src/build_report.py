"""Build the Word submission report from the project's Markdown docs and results.

The Markdown files in docs/ and data/ are the single source of content; this script converts
them (headings, paragraphs, lists, tables, quotes, images, code) into one .docx file, and adds
live model output and the Git commit history as evidence.

Run:  python src/build_report.py
"""
import re
import subprocess
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from predict import predict

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "report" / "Kimuntu_AI_Week1_Report.docx"
DISCLAIMER = ("Kimuntu AI supports user understanding of legal situations. It does not provide legal "
              "advice and does not replace a qualified legal professional.")
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\)|<https?://[^>]+>)")

# New queries, not taken from the dataset (same as the notebook demo)
DEMO_QUERIES = [
    "The landlord is refusing to give back the deposit I paid when I moved in.",
    "I have not received my wages from my employer since July.",
    "What do I need to register my new business?",
    "My contract of employment was terminated without notice.",
    "Who will win the election next year?",
    "I have a problem, can you help?",
]


def add_inline(paragraph, text, bold=False):
    """Add text with **bold**, *italic*, `code` and [link](url) formatting."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            add_inline(paragraph, part[2:-2], bold=True)
        elif part.startswith("`[PLACEHOLDER"):              # make placeholders hard to miss
            run = paragraph.add_run(part[1:-1])
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.bold = bold
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.bold = bold
        elif part.startswith("<http"):
            paragraph.add_run(part[1:-1]).bold = bold
        elif part.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            add_inline(paragraph, label if not url.startswith("http") else f"{label} ({url})", bold)
        elif part.startswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic, run.bold = True, bold
        else:
            paragraph.add_run(part).bold = bold


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_column_widths(table, rows):
    """Give each column a width proportional to its longest text (within limits),
    and never narrower than its longest single word, so words are not broken."""
    weights = []
    for c in range(len(table.columns)):
        texts = [r[c] for r in rows if c < len(r)]
        longest_text = min(max(len(t) for t in texts), 45)
        longest_word = max(len(w) for t in texts for w in (t.split() or [""]))
        weights.append(max(longest_text, longest_word * 2.3) + 4)
    total_cm = 16.6
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, weights):
            cell.width = Cm(total_cm * w / sum(weights))


def add_table(doc, rows):
    header, body = rows[0], [r for r in rows[1:] if not all(set(c) <= set("-: ") for c in r)]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, header):
        cell.paragraphs[0].text = ""
        add_inline(cell.paragraphs[0], text, bold=True)
        shade(cell, "D9E2F3")
    for row in body:
        cells = table.add_row().cells
        for cell, text in zip(cells, row + [""] * (len(header) - len(row))):
            add_inline(cell.paragraphs[0], text)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    set_column_widths(table, [header] + body)
    doc.add_paragraph()


def add_markdown(doc, path, skip_title=True, heading_offset=0):
    """Convert a Markdown file into the document."""
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("```"):                      # code block
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name, run.font.size = "Consolas", Pt(8.5)
            i += 1
            continue
        if stripped.startswith("|"):                         # table
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue
        m = re.match(r"(#+)\s+(.*)", stripped)
        if m:                                                # heading
            level = len(m.group(1))
            if not (level == 1 and skip_title):
                # drop the docs' own "3." numbering so it does not clash with the report's sections
                text = re.sub(r"^\d+\.\s+", "", m.group(2).replace("`", ""))
                doc.add_heading(text, level=min(level + heading_offset, 4))
            i += 1
            continue
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if m:                                                # image
            img = (path.parent / m.group(2)).resolve()
            doc.add_picture(str(img), width=Cm(15))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph(m.group(1))
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
            i += 1
            continue
        if stripped.startswith(">"):                         # block quote
            text = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                text.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, " ".join(text))
            continue
        m = re.match(r"(\s*)([-*]|\d+\.)\s+(.*)", line)
        if m:                                                # list item (with continuation lines)
            indent = len(m.group(1)) // 2
            text = m.group(3)
            i += 1
            while (i < len(lines) and lines[i].strip() and not re.match(r"\s*([-*]|\d+\.)\s+", lines[i])
                   and lines[i].startswith(" ") and not lines[i].strip().startswith("|")):
                text += " " + lines[i].strip()
                i += 1
            if m.group(2)[0].isdigit():
                # keep the Markdown's own numbers; Word's List Number would continue across the document
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9 + 0.6 * indent)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.add_run(f"{m.group(2)} ")
            else:
                p = doc.add_paragraph(style="List Bullet 2" if indent else "List Bullet")
            add_inline(p, text)
            continue
        text = [stripped]                                    # paragraph
        i += 1
        while (i < len(lines) and lines[i].strip() and not re.match(r"\s*(#|\||>|!\[|```|[-*]\s|\d+\.\s)",
                                                                      lines[i])):
            text.append(lines[i].strip())
            i += 1
        add_inline(doc.add_paragraph(), " ".join(text))


def add_model_output(doc):
    doc.add_paragraph(
        "The baseline is implemented in src/train_baseline.py and, equivalently, in "
        "notebooks/baseline_classifier.ipynb (Google Colab compatible). It is a scikit-learn Pipeline: "
        "TfidfVectorizer followed by a classifier. Three model types were compared on the validation set "
        "(majority-class dummy, Multinomial Naive Bayes and Logistic Regression, 33 configurations in total) "
        "and the configuration with the best validation macro-F1 was kept. The model is saved to "
        "models/baseline_classifier.joblib and called through predict() in src/predict.py.")
    doc.add_heading("Why these choices", level=3)
    for text in [
        "TF-IDF turns each query into word weights: frequent-but-generic words get low weight, distinctive "
        "legal words (landlord, probate, dismissal) get high weight. It is simple, fast and explainable.",
        "Logistic Regression and Naive Bayes are the methods recommended in the brief, train in seconds and "
        "can be inspected (top words per class), following the lecturer note's guidance to start with "
        "interpretable baselines in scikit-learn.",
        "class_weight='balanced' makes mistakes on small classes cost more, to counter class imbalance.",
        "The vectoriser sits inside the pipeline, so it is fitted on training data only (no leakage); "
        "the test set was used once, after model selection.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("Model output: predict() on new example queries", level=3)
    doc.add_paragraph("These queries were written for this demonstration and are not in the dataset. "
                      "The output below was produced by the saved model when this report was built.")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Query", "Predicted category", "Confidence", "Needs clarification"]):
        cell.paragraphs[0].add_run(text).bold = True
        shade(cell, "D9E2F3")
    rows = [["Query", "Predicted category", "Confidence", "Needs clarification"]]
    for q in DEMO_QUERIES:
        r = predict(q)
        rows.append([q, r["category"], f"{r['confidence']:.4f}", "yes" if r["needs_clarification"] else "no"])
        for cell, text in zip(table.add_row().cells, rows[-1]):
            cell.paragraphs[0].add_run(text).font.size = Pt(9)
    set_column_widths(table, rows)
    doc.add_paragraph()
    doc.add_paragraph(f"Threshold used by the safety gate: {predict('test')['threshold']}. "
                      "The fifth query is a real failure case, discussed in Section 7.")


def add_git_history(doc):
    log = subprocess.run(["git", "log", "--reverse", "--format=%h|%ad|%an|%s", "--date=short"],
                         cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip().splitlines()
    doc.add_paragraph("Commit history of the repository at the time this report was built "
                      "(the commit that adds the report itself comes after). Run 'git log' for the current history.")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    rows = [["Commit", "Date", "Author", "Message"]]
    for cell, text in zip(table.rows[0].cells, rows[0]):
        cell.paragraphs[0].add_run(text).bold = True
        shade(cell, "D9E2F3")
    for line in log:
        rows.append(line.split("|", 3))
        for cell, text in zip(table.add_row().cells, rows[-1]):
            cell.paragraphs[0].add_run(text).font.size = Pt(8.5)
    set_column_widths(table, rows)
    doc.add_paragraph()
    doc.add_paragraph("All commits carry a 'Co-Authored-By: Claude' line, recording the use of an AI assistant.")


def split_interpretation(tmp_dir):
    """Split docs/interpretation.md into results (sections 1-2) and interpretation (sections 3-5)."""
    text = (ROOT / "docs" / "interpretation.md").read_text(encoding="utf-8")
    head, rest = text.split("## 3.", 1)
    results, interp = tmp_dir / "results_part.md", tmp_dir / "interpretation_part.md"
    results.write_text(head, encoding="utf-8")
    interp.write_text("## 3." + rest, encoding="utf-8")
    return results, interp


def main():
    doc = Document()
    zoom = doc.settings.element.find(qn("w:zoom"))   # default template omits the required attribute
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.2)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_heading("Kimuntu AI", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Week 1: System Design and Baseline Legal Query Classifier")
    run.bold, run.font.size = True, Pt(16)
    for label, value in [("Course", "AI and Expert Systems in Action"),
                         ("Industry partner", "Kimuntu Power Inc. (Riipen industry project)"),
                         ("Lecturer", "Dr. Chinwe Peace Igiri"),
                         ("Student name", "[PLACEHOLDER]"),
                         ("Assigned role", "[PLACEHOLDER]"),
                         ("Date", date.today().isoformat())]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"{label}: ").bold = True
        run = p.add_run(value)
        if value.startswith("[PLACEHOLDER"):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    note = doc.add_paragraph(style="Intense Quote")
    note.add_run("Disclaimer: ").bold = True
    note.add_run(DISCLAIMER)

    doc.add_heading("Contents", level=1)
    for item in ["1. Problem and scope statement", "2. Proposed legal text classification categories",
                 "3. System architecture", "4. Dataset and preprocessing", "5. Baseline text classification model",
                 "6. Model evaluation results", "7. Interpretation, limitations and Week 2 improvements",
                 "8. Group output and individual contribution", "Appendix A. How to reproduce",
                 "Appendix B. References and data sources"]:
        doc.add_paragraph(item)
    doc.add_page_break()

    tmp = ROOT / "report" / "_tmp"
    tmp.mkdir(parents=True, exist_ok=True)
    results_md, interp_md = split_interpretation(tmp)

    sections = [
        ("1. Problem and scope statement", ROOT / "docs" / "problem_scope.md"),
        ("2. Proposed legal text classification categories", ROOT / "docs" / "categories.md"),
        ("3. System architecture", ROOT / "docs" / "architecture.md"),
        ("4. Dataset and preprocessing", ROOT / "data" / "README.md"),
        ("5. Baseline text classification model", None),
        ("6. Model evaluation results", results_md),
        ("7. Interpretation, limitations and Week 2 improvements", interp_md),
    ]
    for heading, md in sections:
        doc.add_heading(heading, level=1)
        if md is None:
            add_model_output(doc)
        else:
            # results_part.md references the confusion matrix relative to docs/
            add_markdown(doc, md if md.parent != tmp else _relocate(md))
        doc.add_page_break()

    doc.add_heading("8. Group output and individual contribution", level=1)
    doc.add_paragraph("This report and the repository it describes form the complete group output for "
                      "Week 1. The individual statement below must be completed by the student.")
    add_markdown(doc, ROOT / "docs" / "individual_contribution.md")
    doc.add_heading("Team role allocation", level=2)
    add_markdown(doc, ROOT / "docs" / "team_roles.md")
    doc.add_heading("Evidence: Git commit history", level=2)
    add_git_history(doc)
    doc.add_page_break()

    doc.add_heading("Appendix A. How to reproduce", level=1)
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    start = readme.index("## How to run locally")
    end = readme.index("## Data sources and licences")
    (tmp / "reproduce.md").write_text(readme[start:end], encoding="utf-8")
    add_markdown(doc, tmp / "reproduce.md", heading_offset=0)

    doc.add_heading("Appendix B. References and data sources", level=1)
    for ref in [
        "Moslem, Y. (2023). Law-StackExchange Dataset. Hugging Face. "
        "https://huggingface.co/datasets/ymoslem/Law-StackExchange. DOI: 10.57967/hf/4792. Licence: CC BY-SA 4.0.",
        "Law Stack Exchange contributors. Questions licensed CC BY-SA 3.0 / 4.0. https://law.stackexchange.com",
        "Uganda Legal Information Institute (ULII), Law Reporting Unit of the Uganda Judiciary. https://ulii.org/ "
        "(candidate retrieval source for later weeks, pending Kimuntu approval).",
        "Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning "
        "Research, 12, 2825-2830.",
        "Assessment brief (AI.pdf) and Kimuntu AI Riipen Project Discussion Note (lecturer guide), course documents.",
        "Synthetic queries: written for this project by an AI assistant (Claude, Anthropic); not real user data.",
    ]:
        doc.add_paragraph(ref, style="List Bullet")

    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    for f in tmp.iterdir():
        f.unlink()
    tmp.rmdir()
    print(f"Saved {OUT}")


def _relocate(md):
    """Rewrite image paths in a temporary file so they resolve from report/_tmp/."""
    text = md.read_text(encoding="utf-8").replace("](../results/", "](../../results/")
    md.write_text(text, encoding="utf-8")
    return md


if __name__ == "__main__":
    main()
